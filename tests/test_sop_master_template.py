"""
tests/test_sop_master_template.py — Yuktav Master SOP Template.

Covers the master SOP structure (product-owner reference: "SOP Template.docx")
implemented on top of the existing Document Control architecture:

  1. prompts/qms_document_prompt.py::build_draft_prompt() applies
     DEFAULT_SOP_STRUCTURE for doc_type=='SOP' with no explicit template.
  2. database.py seeds/refreshes the same structure onto the platform-wide
     "Standard Operating Procedure (Default)" qms_document_templates row.
  3. services/docx_generator.py's SOP-specific build path (_build_sop):
     section order, header/footer, nested Procedure sub-headings, and the
     Revision History / Approvals tables at the end of the document.
  4. routes/qms_documents.py export_docx() passes the document's own
     doc_number/version/effective_date/status/owner/reviewer/approver
     through to the DOCX (previously dropped on the floor).

Does not touch the legacy /validation/generate SOP path or any other doc_type
— those are out of scope for this change and are covered by their own tests.
"""

import io

from docx import Document as DocxDocument

from pharmagpt.prompts.qms_document_prompt import DEFAULT_SOP_STRUCTURE, build_draft_prompt
from pharmagpt.services.docx_generator import build_document_data, generate_docx


# ── Prompt structure ──────────────────────────────────────────────────────────

def test_default_sop_structure_matches_master_section_numbering():
    headings = [s["heading"] for s in DEFAULT_SOP_STRUCTURE]
    assert headings == [
        "2. Purpose",
        "3. Scope",
        "4. Responsibility",
        "5. References",
        "6. Abbreviation/Definition",
        "7. Safety",
        "8. Procedure",
        "9. FORM/ANNEXURE",
    ]


def test_build_draft_prompt_applies_master_structure_with_no_explicit_template():
    prompt = build_draft_prompt({"title": "Cleaning of Tablet Press", "doc_type": "SOP"})
    for heading in ["## 2. Purpose", "## 3. Scope", "## 4. Responsibility", "## 5. References",
                     "## 6. Abbreviation/Definition", "## 7. Safety", "## 8. Procedure",
                     "## 9. FORM/ANNEXURE"]:
        assert heading in prompt
    # Table-column hints for the tabular sections.
    assert "Department/User | Responsibility" in prompt
    assert "SOP No. | SOP Title" in prompt
    assert "Abbreviation/Definition | Full Form/Meaning" in prompt
    assert "Form No. | Title" in prompt
    # Procedure stays flexible — nesting guidance present, no forced sub-headings.
    assert "8.1" in prompt and "8.1.1" in prompt
    # Content/Revision History/Approvals are DOCX-native, not requested from the AI.
    assert "Content" not in prompt.split("DOCUMENT INFORMATION:")[1].split("## 2. Purpose")[0]
    assert "Revision History" not in prompt
    assert "Approvals" not in prompt


def test_build_draft_prompt_explicit_template_still_overrides_default():
    custom = {"name": "Custom Co. SOP", "structure": [{"heading": "1. Custom Heading", "sub_headings": []}]}
    prompt = build_draft_prompt({"title": "X", "doc_type": "SOP"}, template=custom)
    assert "1. Custom Heading" in prompt
    assert "2. Purpose" not in prompt


# ── Seeded default template row (database.py) ────────────────────────────────

def test_seeded_default_sop_template_matches_master_structure(client):
    templates = client.get("/qms/documents/templates?doc_type=SOP").get_json()
    default_template = next(t for t in templates if t["name"] == "Standard Operating Procedure (Default)")
    assert [s["heading"] for s in default_template["structure"]] == \
        [s["heading"] for s in DEFAULT_SOP_STRUCTURE]


# ── DOCX assembly (services/docx_generator.py _build_sop) ────────────────────

_SAMPLE_CONTENT = """# Cleaning of Tablet Press

## 2. Purpose
This SOP describes the cleaning procedure.

## 3. Scope
Applies to the Production department.

## 4. Responsibility
| Department/User | Responsibility |
|---|---|
| Operator | Perform cleaning |

## 5. References
| SOP No. | SOP Title |
|---|---|
| SOP-00099 | Equipment Qualification |

## 6. Abbreviation/Definition
| Abbreviation/Definition | Full Form/Meaning |
|---|---|
| IPA | Isopropyl Alcohol |

## 7. Safety
Wear gloves and safety glasses.

## 8. Procedure
### 8.1 Operation
#### 8.1.1 Switch on the machine
#### 8.1.2 Set parameters
### 8.2 Cleaning
#### 8.2.1 Disassemble parts
#### 8.2.2 Clean with IPA

## 9. FORM/ANNEXURE
| Form No. | Title |
|---|---|
| FRM-001 | Cleaning Log |
"""


def _build_sop_docx(**form_overrides):
    form_data = {
        "title": "Cleaning of Tablet Press",
        "department": "Production",
        "protocol_number": "SOP-00123",
        "revision_number": "1.0",
        "effective_date": "01-Sep-2026",
        "document_status": "Draft",
        "company_name": "Yuktav Pharma Ltd",
        "logo_path": None,
        "prepared_by": "A. Kumar",
        "reviewed_by": "B. Singh",
        "approved_by": "C. Rao",
        **form_overrides,
    }
    data = build_document_data(form_data, "SOP", form_data["title"], _SAMPLE_CONTENT, project_id=0)
    docx_bytes = generate_docx(data)
    return DocxDocument(io.BytesIO(docx_bytes))


def test_sop_docx_section_order():
    doc = _build_sop_docx()
    headings = [p.text for p in doc.paragraphs
                if p.style.name.startswith("Heading") or p.text.strip() == "1. CONTENT"]
    # Strip the bare document-title Heading-1 paragraph that precedes "2. Purpose".
    ordered = [h for h in headings if h != "Cleaning of Tablet Press"]
    assert ordered == [
        "1. CONTENT",
        "2. Purpose",
        "3. Scope",
        "4. Responsibility",
        "5. References",
        "6. Abbreviation/Definition",
        "7. Safety",
        "8. Procedure",
        "8.1 Operation",
        "8.1.1 Switch on the machine",
        "8.1.2 Set parameters",
        "8.2 Cleaning",
        "8.2.1 Disassemble parts",
        "8.2.2 Clean with IPA",
        "9. FORM/ANNEXURE",
        "10. REVISION HISTORY",
        "11. APPROVALS",
    ]


def test_sop_docx_procedure_subsections_use_nested_heading_styles():
    doc = _build_sop_docx()
    by_text = {p.text: p.style.name for p in doc.paragraphs if p.text.strip()}
    assert by_text["8. Procedure"] == "Heading 1"
    assert by_text["8.1 Operation"] == "Heading 2"
    assert by_text["8.1.1 Switch on the machine"] == "Heading 3"
    assert by_text["8.2 Cleaning"] == "Heading 2"


def test_sop_docx_no_separate_cover_page():
    doc = _build_sop_docx()
    # Header/footer must be visible from page 1 (no cover-page suppression).
    assert doc.sections[0].different_first_page_header_footer is False


def test_sop_docx_header_is_dynamic_from_company_and_document_data():
    doc = _build_sop_docx()
    header_tables = doc.sections[0].header.tables
    assert len(header_tables) == 1
    rows = [[c.text.strip() for c in row.cells] for row in header_tables[0].rows]
    assert rows[0][1] == "Cleaning of Tablet Press"   # SOP title, merged across cols 1-2
    assert rows[0][2] == "Cleaning of Tablet Press"
    assert rows[1][0] == "Yuktav Pharma Ltd"           # company name
    assert rows[2][0] == "Yuktav Pharma Ltd"           # vertically merged
    assert rows[1][1] == "SOP NO: SOP-00123"
    assert rows[1][2] == "Version No: 1.0"
    assert rows[2][1] == "Effective Date: 01-Sep-2026"

    xml = doc.sections[0].header._element.xml
    assert "PAGE" in xml and "NUMPAGES" in xml   # native Word page-numbering fields


def test_sop_docx_header_shows_placeholder_when_no_branding_supplied():
    doc = _build_sop_docx(company_name="", logo_path=None)
    header_tables = doc.sections[0].header.tables
    rows = [[c.text.strip() for c in row.cells] for row in header_tables[0].rows]
    assert rows[0][0] == "[Company Logo]"
    assert rows[1][0] == "Company Name"


def test_sop_docx_footer_shows_only_sop_number_and_version():
    doc = _build_sop_docx()
    footer_text = doc.sections[0].footer.paragraphs[0].text
    assert footer_text == "SOP-00123\tVersion No: 1.0"
    assert "PharmaGPT" not in footer_text
    assert "CONTROLLED DOCUMENT" not in footer_text


def test_sop_docx_revision_history_at_end_with_required_columns():
    doc = _build_sop_docx()
    tables = doc.tables
    rev_table = next(t for t in tables if [c.text.strip() for c in t.rows[0].cells] ==
                      ["Revision No.", "Author", "Change History"])
    assert rev_table.rows[1].cells[0].text.strip() == "1.0"
    assert rev_table.rows[1].cells[1].text.strip() == "A. Kumar"


def test_sop_docx_approvals_at_end_with_required_columns():
    doc = _build_sop_docx()
    approvals_table = next(t for t in doc.tables if [c.text.strip() for c in t.rows[0].cells] ==
                            ["Name", "Department"])
    assert approvals_table is not None


def test_sop_docx_approvals_accepts_explicit_signatories():
    doc = _build_sop_docx(signatories=[{"name": "A. Kumar", "department": "Production"},
                                        {"name": "B. Singh", "department": "Quality Assurance"}])
    approvals_table = next(t for t in doc.tables if [c.text.strip() for c in t.rows[0].cells] ==
                            ["Name", "Department"])
    rows = [[c.text.strip() for c in row.cells] for row in approvals_table.rows[1:]]
    assert rows == [["A. Kumar", "Production"], ["B. Singh", "Quality Assurance"]]


def test_sop_docx_form_annexure_table_from_content_not_hardcoded_equipment_annexures():
    doc = _build_sop_docx()
    form_table = next(t for t in doc.tables if [c.text.strip() for c in t.rows[0].cells] ==
                       ["Form No.", "Title"])
    assert form_table.rows[1].cells[0].text.strip() == "FRM-001"
    # The generic qualification-protocol annexure placeholders must not appear.
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Calibration Certificates" not in body_text
    assert "Equipment Drawings" not in body_text


def test_non_sop_doc_type_still_uses_generic_cover_page_pipeline():
    """Regression guard: doc_type != 'SOP' must be completely unaffected by
    the SOP-specific build path added in this change."""
    form_data = {"title": "HPLC OQ Protocol", "company_name": "Yuktav Pharma Ltd"}
    data = build_document_data(form_data, "OQ", "HPLC OQ Protocol", "## Purpose\nText.", project_id=0)
    doc = DocxDocument(io.BytesIO(generate_docx(data)))
    assert doc.sections[0].different_first_page_header_footer is True
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "TABLE OF CONTENTS" in body_text
    assert "CONFIDENTIAL — CONTROLLED DOCUMENT" in body_text


# ── routes/qms_documents.py export_docx() field pass-through ─────────────────

def test_export_docx_route_passes_document_metadata_into_docx(client):
    doc = client.post("/qms/documents", json={
        "title": "Operation of Tablet Press", "doc_type": "SOP", "department": "Production",
        "doc_number": "SOP-00456", "effective_date": "15-Aug-2026",
        "owner": "A. Kumar", "content": "## 2. Purpose\nText.\n",
    }).get_json()
    # "version" is system-managed by the version ledger (services/document_
    # versioning.py), not a free-form field a caller can set at creation —
    # read back whatever the system actually assigned rather than assuming one.
    assert doc["doc_number"] == "SOP-00456"

    resp = client.post(f"/qms/documents/{doc['id']}/export/docx",
                        json={"company_name": "Yuktav Pharma Ltd"})
    assert resp.status_code == 200
    assert resp.mimetype == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    docx_doc = DocxDocument(io.BytesIO(resp.data))
    header_tables = docx_doc.sections[0].header.tables
    rows = [[c.text.strip() for c in row.cells] for row in header_tables[0].rows]
    assert rows[1][1] == "SOP NO: SOP-00456"
    assert rows[1][2] == f"Version No: {doc['version']}"
    assert rows[2][1] == "Effective Date: 15-Aug-2026"
    assert rows[1][0] == "Yuktav Pharma Ltd"


# ── Company Branding bridge (services/company_branding.py) ───────────────────

def test_get_company_name_for_docx_returns_empty_for_falsy_company_id():
    """super_admin outside Assume Company Context has company_id=None
    (pharmagpt/auth/context.py) — must resolve to "" without attempting any
    network call, not raise."""
    from pharmagpt.services.company_branding import get_company_name_for_docx
    assert get_company_name_for_docx(None) == ""
    assert get_company_name_for_docx("") == ""


def test_get_company_name_for_docx_returns_empty_with_no_request_context():
    """Outside a Flask request (no Authorization header to extract), must
    fail closed to "" rather than raise."""
    from pharmagpt.services.company_branding import get_company_name_for_docx
    assert get_company_name_for_docx("some-company-id") == ""


def _make_sop_document(client, **overrides):
    payload = {
        "title": "Cleaning of Tablet Press", "doc_type": "SOP", "department": "Production",
        "doc_number": "SOP-00789", "content": "## 2. Purpose\nText.\n",
    }
    payload.update(overrides)
    return client.post("/qms/documents", json=payload).get_json()


def test_export_docx_auto_populates_company_name_when_not_overridden(client, monkeypatch):
    captured = {}

    def _fake_lookup(company_id):
        captured["company_id"] = company_id
        return "Auto-Resolved Pharma Ltd"

    monkeypatch.setattr("pharmagpt.routes.qms_documents.get_company_name_for_docx", _fake_lookup)

    doc = _make_sop_document(client)
    resp = client.post(f"/qms/documents/{doc['id']}/export/docx", json={})
    assert resp.status_code == 200

    docx_doc = DocxDocument(io.BytesIO(resp.data))
    rows = [[c.text.strip() for c in row.cells] for row in docx_doc.sections[0].header.tables[0].rows]
    assert rows[1][0] == "Auto-Resolved Pharma Ltd"
    # Confirms the lookup used the server-derived tenant company_id, not
    # anything from the request body — see the tenant-isolation test below.
    assert captured["company_id"] == doc["company_id"]


def test_export_docx_explicit_company_name_overrides_auto_lookup(client, monkeypatch):
    monkeypatch.setattr(
        "pharmagpt.routes.qms_documents.get_company_name_for_docx",
        lambda company_id: "Should Not Be Used",
    )
    doc = _make_sop_document(client)
    resp = client.post(f"/qms/documents/{doc['id']}/export/docx",
                        json={"company_name": "Explicit Override Ltd"})
    docx_doc = DocxDocument(io.BytesIO(resp.data))
    rows = [[c.text.strip() for c in row.cells] for row in docx_doc.sections[0].header.tables[0].rows]
    assert rows[1][0] == "Explicit Override Ltd"


def test_download_template_docx_also_auto_populates_company_name(client, monkeypatch):
    monkeypatch.setattr(
        "pharmagpt.routes.qms_documents.get_company_name_for_docx",
        lambda company_id: "Template Preview Co",
    )
    templates = client.get("/qms/documents/templates?doc_type=SOP").get_json()
    default_template = next(t for t in templates if t["name"] == "Standard Operating Procedure (Default)")
    resp = client.get(f"/qms/documents/templates/{default_template['id']}/download")
    assert resp.status_code == 200
    docx_doc = DocxDocument(io.BytesIO(resp.data))
    rows = [[c.text.strip() for c in row.cells] for row in docx_doc.sections[0].header.tables[0].rows]
    assert rows[1][0] == "Template Preview Co"


# ── Security: tenant isolation / no client-controlled company_id or path ────

def test_export_docx_ignores_client_supplied_company_id_for_branding_lookup(client, monkeypatch):
    """A request body cannot redirect the branding lookup to another
    tenant's company — export_docx must always pass g.tenant.company_id
    (server-derived), never anything read from the request body."""
    captured = {}
    monkeypatch.setattr(
        "pharmagpt.routes.qms_documents.get_company_name_for_docx",
        lambda company_id: captured.setdefault("company_id", company_id) or "irrelevant",
    )
    doc = _make_sop_document(client)
    resp = client.post(f"/qms/documents/{doc['id']}/export/docx",
                        json={"company_id": "attacker-controlled-other-company-id"})
    assert resp.status_code == 200
    assert captured["company_id"] == doc["company_id"]
    assert captured["company_id"] != "attacker-controlled-other-company-id"


def test_export_docx_no_longer_accepts_client_supplied_logo_path(client, monkeypatch, tmp_path):
    """Regression test for the local-file-read primitive removed in this
    change: a request-supplied logo_path must be completely ignored — the
    server must never attempt to open a caller-chosen filesystem path and
    embed its contents into the exported document."""
    monkeypatch.setattr("pharmagpt.routes.qms_documents.get_company_name_for_docx", lambda company_id: "")

    # A real, readable file the app process CAN open — if the (removed)
    # code path were still active, this would successfully "load" as an
    # attempted picture and prove the vulnerability is still present.
    sentinel = tmp_path / "sentinel.txt"
    sentinel.write_text("not an image, should never be read by export_docx")

    doc = _make_sop_document(client)
    resp = client.post(f"/qms/documents/{doc['id']}/export/docx",
                        json={"logo_path": str(sentinel)})
    assert resp.status_code == 200   # must not error attempting to open/parse the file as an image

    docx_doc = DocxDocument(io.BytesIO(resp.data))
    rows = [[c.text.strip() for c in row.cells] for row in docx_doc.sections[0].header.tables[0].rows]
    assert rows[0][0] == "[Company Logo]"   # falls back to the placeholder, never the sentinel path


def test_export_docx_branding_lookup_failure_does_not_block_export(client, monkeypatch):
    """A branding-lookup failure (network error, Supabase misconfiguration,
    etc.) must degrade to no company name, never fail the export itself."""
    def _boom(company_id):
        raise RuntimeError("simulated Supabase outage")
    # get_company_name_for_docx itself already catches and swallows this
    # (see its own docstring/implementation) — this exercises that contract
    # directly rather than depending on export_docx to also catch it.
    from pharmagpt.services import company_branding
    import pharmagpt.routes.qms_documents as qms_documents_route
    real_get_authenticated_client = company_branding.get_authenticated_client
    monkeypatch.setattr(company_branding, "get_authenticated_client",
                         lambda token: (_ for _ in ()).throw(RuntimeError("simulated Supabase outage")))
    monkeypatch.setattr(company_branding, "extract_bearer_token", lambda: "fake-token")
    monkeypatch.setattr(qms_documents_route, "get_company_name_for_docx",
                         company_branding.get_company_name_for_docx)

    doc = _make_sop_document(client)
    resp = client.post(f"/qms/documents/{doc['id']}/export/docx", json={})
    assert resp.status_code == 200   # export still succeeds
    docx_doc = DocxDocument(io.BytesIO(resp.data))
    rows = [[c.text.strip() for c in row.cells] for row in docx_doc.sections[0].header.tables[0].rows]
    assert rows[1][0] == "Company Name"   # placeholder, not a crash
